import base64
import io
import logging

from odoo import models, fields, api
from markupsafe import Markup

_logger = logging.getLogger(__name__)


class DocumentPreviewWizard(models.TransientModel):
    _name = 'commercial.document.preview.wizard'
    _description = 'Aperçu du document'

    document_id = fields.Many2one('commercial.document.partner', string="Document")
    fichier_type = fields.Char(related='document_id.fichier_type')
    fichier_name = fields.Char(related='document_id.fichier_name')
    preview_html = fields.Html(
        string="Aperçu", compute='_compute_preview_html', sanitize=False)

    def _is_public_server(self):
        """Vérifie si le serveur est accessible depuis internet."""
        base_url = self.env['ir.config_parameter'].sudo().get_param('web.base.url') or ''
        return 'localhost' not in base_url and '127.0.0.1' not in base_url

    def _get_file_bytes(self):
        """Retourne les bytes bruts du fichier depuis ir.attachment."""
        attachment = self.env['ir.attachment'].sudo().with_context(bin_size=False).search([
            ('res_model', '=', 'commercial.document.partner'),
            ('res_id', '=', self.document_id.id),
            ('res_field', '=', 'fichier'),
        ], limit=1)
        if attachment:
            b64 = attachment.with_context(bin_size=False).datas
            if b64:
                if isinstance(b64, str):
                    b64 = b64.encode('utf-8')
                return base64.b64decode(b64)
        return None

    def _render_office_online(self):
        """Iframe Office Online Viewer (serveur public uniquement)."""
        doc = self.document_id
        token = doc.access_token
        if not token:
            token = doc._generate_access_token()

        base_url = self.env['ir.config_parameter'].sudo().get_param('web.base.url')
        public_url = '%s/commercial/document/public/%s?token=%s' % (
            base_url, doc.id, token)
        viewer_url = 'https://view.officeapps.live.com/op/embed.aspx?src=%s' % public_url

        return Markup(
            '<iframe src="%s" '
            'style="width:100%%;height:80vh;border:none;" '
            'frameborder="0"></iframe>'
        ) % viewer_url

    def _render_xlsx(self):
        """Convertit un fichier XLSX en tableau HTML."""
        import openpyxl
        raw = self._get_file_bytes()
        if not raw:
            return Markup('<p style="text-align:center;color:#888;">Fichier introuvable</p>')
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        html_parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            html_parts.append(
                '<h5 style="margin:15px 0 8px;color:#333;">%s</h5>' % sheet)
            html_parts.append(
                '<div style="overflow-x:auto;margin-bottom:15px;">'
                '<table style="border-collapse:collapse;width:100%%;font-size:12px;">')
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                tag = 'th' if i == 0 else 'td'
                bg = 'background:#e9ecef;font-weight:600;' if i == 0 else ''
                html_parts.append('<tr>')
                for cell in row:
                    val = '' if cell is None else str(cell)
                    html_parts.append(
                        '<%s style="border:1px solid #dee2e6;padding:6px 10px;%s">%s</%s>'
                        % (tag, bg, val, tag))
                html_parts.append('</tr>')
            html_parts.append('</table></div>')
        wb.close()
        return Markup(''.join(html_parts))

    def _render_docx(self):
        """Convertit un fichier DOCX en HTML."""
        from docx import Document
        raw = self._get_file_bytes()
        if not raw:
            return Markup('<p style="text-align:center;color:#888;">Fichier introuvable</p>')
        doc = Document(io.BytesIO(raw))
        html_parts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                html_parts.append('<br/>')
                continue
            style = para.style.name or ''
            if 'Heading 1' in style:
                html_parts.append('<h3 style="color:#333;margin:12px 0 6px;">%s</h3>' % text)
            elif 'Heading 2' in style:
                html_parts.append('<h4 style="color:#444;margin:10px 0 5px;">%s</h4>' % text)
            elif 'Heading' in style:
                html_parts.append('<h5 style="color:#555;margin:8px 0 4px;">%s</h5>' % text)
            else:
                html_parts.append(
                    '<p style="margin:3px 0;line-height:1.6;font-size:13px;">%s</p>' % text)
        for table in doc.tables:
            html_parts.append(
                '<table style="border-collapse:collapse;width:100%%;font-size:12px;margin:10px 0;">')
            for i, row in enumerate(table.rows):
                tag = 'th' if i == 0 else 'td'
                bg = 'background:#e9ecef;font-weight:600;' if i == 0 else ''
                html_parts.append('<tr>')
                for cell in row.cells:
                    html_parts.append(
                        '<%s style="border:1px solid #dee2e6;padding:6px 10px;%s">%s</%s>'
                        % (tag, bg, cell.text, tag))
                html_parts.append('</tr>')
            html_parts.append('</table>')
        return Markup(''.join(html_parts)) if html_parts else Markup(
            '<p style="color:#888;text-align:center;">Document vide</p>')

    @api.depends('document_id')
    def _compute_preview_html(self):
        for rec in self:
            if not rec.document_id:
                rec.preview_html = Markup(
                    '<p style="text-align:center;color:#888;padding:40px;">'
                    'Aucun fichier joint</p>')
                continue

            ftype = (rec.fichier_type or '').upper()
            doc_id = rec.document_id.id

            try:
                if ftype == 'PDF':
                    url = '/commercial/document/preview/%s' % doc_id
                    rec.preview_html = Markup(
                        '<iframe src="%s" '
                        'style="width:100%%;height:80vh;border:none;"></iframe>'
                    ) % url

                elif ftype in ('PNG', 'JPG', 'JPEG', 'GIF', 'BMP', 'WEBP'):
                    url = '/commercial/document/preview/%s' % doc_id
                    rec.preview_html = Markup(
                        '<div style="text-align:center;padding:10px;">'
                        '<img src="%s" style="max-width:100%%;max-height:75vh;"/>'
                        '</div>'
                    ) % url

                elif ftype in ('DOCX', 'DOC', 'XLSX', 'XLS'):
                    if rec._is_public_server():
                        # En ligne : Office Online Viewer
                        rec.preview_html = rec._render_office_online()
                    else:
                        # En local : rendu HTML
                        if ftype in ('XLSX', 'XLS'):
                            rec.preview_html = rec._render_xlsx()
                        else:
                            rec.preview_html = rec._render_docx()

                else:
                    rec.preview_html = Markup(
                        '<div style="text-align:center;padding:40px;">'
                        '<i class="fa fa-file-o" style="font-size:64px;color:#888;"></i>'
                        '<p style="margin-top:15px;font-size:14px;color:#555;">'
                        'Aperçu non disponible pour .%s</p></div>'
                    ) % ftype

            except Exception as e:
                _logger.exception("Erreur aperçu document %s: %s", doc_id, e)
                rec.preview_html = Markup(
                    '<div style="text-align:center;padding:40px;">'
                    '<i class="fa fa-exclamation-circle" '
                    'style="font-size:48px;color:#dc3545;"></i>'
                    '<p style="margin-top:10px;color:#555;">'
                    'Erreur lors de l\'ouverture du fichier</p>'
                    '<p style="font-size:12px;color:#999;">%s</p></div>'
                ) % str(e)
